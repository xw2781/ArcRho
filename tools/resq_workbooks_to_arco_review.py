"""Copy ResQ-linked workbooks, switch them to the Arco add-in, refresh, and review.

For every workbook under INPUT_FOLDER (subfolders included) this script:

1. copies it to the same relative place under OUTPUT_FOLDER;
2. opens the copy in a private Excel instance and records each ResQ formula
   block's saved values - the figures ResQ produced when the file was last saved;
3. converts the formulas and add-in links to Arco with the Arco add-in's own
   Reset Add-in References routine, which owns the rename rules and handles both
   the 32-bit and 64-bit ResQ add-in paths;
4. runs the add-in's Refresh Workbook and saves the copy;
5. compares every dataset cell by cell with the ResQ values at two decimal places.

The review workbook lands in the Downloads folder. Its Datasets sheet lists every
dataset with links to the dataset in the ResQ original and the Arco copy; the
Differences sheet lists each differing cell; Side by Side shows the ResQ, Arco and
difference grids of each differing dataset.

Run on a PC with Excel, pywin32, openpyxl and an Arco credential:

    py -3.10 tools/resq_workbooks_to_arco_review.py
"""

from __future__ import annotations

import os
import shutil
import stat
import sys
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from win32com.client import DispatchEx

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "python-api" / "migration"))
from resq_migration.engine_parity import (  # noqa: E402  - the ResQ import's two-decimal rule
    IMPORT_ABSOLUTE_TOLERANCE,
    IMPORT_DECIMAL_PLACES,
)

# ---------------------------------------------------------------------------
# Fill these in before running.
INPUT_FOLDER = r""   # folder holding the ResQ-connected workbooks
OUTPUT_FOLDER = r""  # new folder that receives the Arco copies
# ---------------------------------------------------------------------------

ADDIN_PATH = r"E:\ArcRho Server\Excel Add-ins\ArcRho.xlam"
REVIEW_PATH = Path.home() / "Downloads" / f"ResQ_vs_Arco_Review_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
WORKBOOK_SUFFIXES = {".xlsx", ".xlsm", ".xlsb", ".xls"}
SNAPSHOT_SHEET = "_ArcRhoCache"  # the add-in's hidden result store, never compared

XL_MANUAL = -4135
XL_AUTOMATIC = -4105
# pywin32 hands back a cell error as this offset plus Excel's error number.
EXCEL_ERROR_BASE = -2146828288
EXCEL_ERRORS = {
    2000: "#NULL!", 2007: "#DIV/0!", 2015: "#VALUE!", 2023: "#REF!", 2029: "#NAME?",
    2036: "#NUM!", 2042: "#N/A", 2043: "#GETTING_DATA", 2045: "#SPILL!", 2046: "#CONNECT!",
    2047: "#BLOCKED!", 2048: "#UNKNOWN!", 2049: "#FIELD!", 2050: "#CALC!",
}

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="44546A")
DIFF_FILL = PatternFill("solid", fgColor="F8CBAD")
MATCH_FILL = PatternFill("solid", fgColor="E2EFDA")
LINK_FONT = Font(color="0563C1", underline="single")
TITLE_FONT = Font(bold=True, size=12)


@dataclass
class Extent:
    row: int
    col: int
    values: list[list[object]]

    @property
    def rows(self) -> int:
        return len(self.values)

    @property
    def cols(self) -> int:
        return max((len(r) for r in self.values), default=0)

    def value(self, row: int, col: int) -> object:
        r, c = row - self.row, col - self.col
        if 0 <= r < self.rows and 0 <= c < len(self.values[r]):
            return self.values[r][c]
        return None

    @property
    def address(self) -> str:
        return cell_range(self.row, self.col, self.rows, self.cols)


@dataclass
class Dataset:
    sheet: str
    resq_formula: str
    resq: Extent
    arco_formula: str = ""
    arco: Extent | None = None
    status: str = ""
    cells_compared: int = 0
    diffs: list[tuple[int, int, object, object]] = field(default_factory=list)
    max_abs_diff: float | None = None
    side_by_side_row: int | None = None


@dataclass
class BookResult:
    relative: Path
    source: Path
    target: Path
    datasets: list[Dataset] = field(default_factory=list)
    conversion: str = ""
    refresh: str = ""
    error: str = ""


def cell_range(row: int, col: int, rows: int = 1, cols: int = 1) -> str:
    start = f"{get_column_letter(col)}{row}"
    if rows <= 1 and cols <= 1:
        return start
    return f"{start}:{get_column_letter(col + max(cols, 1) - 1)}{row + max(rows, 1) - 1}"


def excel_value(value: object) -> object:
    if isinstance(value, int) and not isinstance(value, bool):
        code = value - EXCEL_ERROR_BASE
        if code in EXCEL_ERRORS:
            return EXCEL_ERRORS[code]
    return value


def read_extent(block) -> tuple[Extent, str]:
    """Return the formula's whole result (legacy array, spill, or single cell) and its formula."""

    anchor = block.Cells(1, 1)
    if anchor.HasArray:
        rng, formula = anchor.CurrentArray, str(anchor.FormulaArray)
    else:
        rng = anchor.SpillingToRange if anchor.HasSpill else anchor
        formula = str(anchor.Formula2)
    raw = rng.Value2
    rows = [list(r) for r in raw] if isinstance(raw, tuple) else [[raw]]
    values = [[excel_value(v) for v in r] for r in rows]
    return Extent(rng.Row, rng.Column, values), formula


def is_missing(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def is_number(value: object) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def same_value(resq: object, arco: object) -> bool:
    if is_missing(resq) or is_missing(arco):
        return is_missing(resq) and is_missing(arco)
    if is_number(resq) and is_number(arco):
        return abs(resq - arco) <= IMPORT_ABSOLUTE_TOLERANCE
    return resq == arco


def compare(dataset: Dataset) -> None:
    a, b = dataset.resq, dataset.arco
    top, left = min(a.row, b.row), min(a.col, b.col)
    bottom = max(a.row + a.rows, b.row + b.rows) - 1
    right = max(a.col + a.cols, b.col + b.cols) - 1
    for row in range(top, bottom + 1):
        for col in range(left, right + 1):
            resq, arco = a.value(row, col), b.value(row, col)
            dataset.cells_compared += 1
            if same_value(resq, arco):
                continue
            dataset.diffs.append((row, col, resq, arco))
            if is_number(resq) and is_number(arco):
                dataset.max_abs_diff = max(dataset.max_abs_diff or 0.0, abs(arco - resq))
    shape_note = "" if (a.rows, a.cols) == (b.rows, b.cols) else " (size changed)"
    dataset.status = ("Different" if dataset.diffs else "Match") + shape_note


def run(excel, name: str, *args):
    return excel.Run(f"'ArcRho.xlam'!{name}", *args)


def process_book(excel, target_path: str, result: BookResult) -> None:
    book = excel.Workbooks.Open(str(result.target), 0, False)
    try:
        # The copy is still byte-identical to the original here, and calculation
        # is manual, so each block still shows the values ResQ last saved. The
        # add-in's own finder returns every candidate block, legacy arrays whole.
        candidates: list[Dataset] = []
        for ws in book.Worksheets:
            if ws.Name == SNAPSHOT_SHEET:
                continue
            blocks = run(excel, "WorkbookRefresh.FindArcRhoFormulaBlocks", ws.UsedRange, False, 0, 15, True)
            for index in range(1, blocks.Count() + 1):
                extent, formula = read_extent(blocks.Item(index))
                candidates.append(Dataset(ws.Name, formula, extent))

        # The same routine as the ribbon's Reset Add-in References > Update workbook.
        result.conversion = str(run(excel, "ReferenceRepair.UpdateWorkbookReferences", book, target_path, "Arco"))
        book.Activate()
        run(excel, "RibbonActions.CalculateWorkbookNoUI")
        result.refresh = str(excel.StatusBar or "").removeprefix(f"Arco [{book.Name}]: ")
        excel.StatusBar = False

        for dataset in candidates:
            anchor = book.Worksheets(dataset.sheet).Cells(dataset.resq.row, dataset.resq.col)
            dataset.arco, dataset.arco_formula = read_extent(anchor)
            if dataset.arco_formula == dataset.resq_formula:
                continue  # not an add-in formula, or one the conversion reports as not updated
            compare(dataset)
            result.datasets.append(dataset)

        # Save in automatic mode so the copy does not reopen in manual calculation.
        excel.Calculation = XL_AUTOMATIC
        # Excel asks "Save with references to unsaved documents?" - a box that
        # ignores DisplayAlerts and hangs the run - when the add-in looks changed.
        for other in excel.Workbooks:
            if other.Name != book.Name:
                other.Saved = True
        book.Save()
        excel.Calculation = XL_MANUAL
    finally:
        book.Close(False)


def convert_all(books: list[BookResult]) -> None:
    excel = DispatchEx("Excel.Application")
    try:
        excel.Visible = False
        excel.DisplayAlerts = False
        excel.AskToUpdateLinks = False
        excel.ScreenUpdating = False
        # Calculation can only be set with a workbook open; the blank one stays open.
        excel.Workbooks.Add()
        excel.Calculation = XL_MANUAL
        # Opening the add-in runs its startup code, which needs events.
        excel.EnableEvents = True
        excel.Workbooks.Open(ADDIN_PATH, 0, True)
        target_path = str(run(excel, "ReferenceRepair.ReferenceTargetPath", 0))
        for index, result in enumerate(books, start=1):
            started = time.monotonic()
            print(f"[{index}/{len(books)}] {result.relative}", flush=True)
            try:
                process_book(excel, target_path, result)
            except Exception as exc:  # one broken workbook must not stop the batch
                result.error = str(exc)
            differing = sum(1 for d in result.datasets if d.diffs)
            print(f"    {len(result.datasets)} dataset(s), {differing} different, "
                  f"refresh: {result.refresh or '-'}{'  ERROR: ' + result.error if result.error else ''} "
                  f"({time.monotonic() - started:.0f}s)", flush=True)
    finally:
        excel.DisplayAlerts = False
        for book in list(excel.Workbooks):
            book.Close(False)
        excel.Quit()


def link(cell, text: str, target: str | None, location: str) -> None:
    cell.value = text
    cell.hyperlink = Hyperlink(ref=cell.coordinate, target=target, location=location)
    cell.font = LINK_FONT


def sheet_location(sheet: str, address: str) -> str:
    return "'" + sheet.replace("'", "''") + "'!" + address


def header(ws, labels: list[str], widths: list[int]) -> None:
    ws.append(labels)
    for index, width in enumerate(widths, start=1):
        cell = ws.cell(1, index)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"


def write_side_by_side(ws, result: BookResult, dataset: Dataset, row: int) -> int:
    a, b = dataset.resq, dataset.arco
    top, left = min(a.row, b.row), min(a.col, b.col)
    rows = max(a.row + a.rows, b.row + b.rows) - top
    cols = max(a.col + a.cols, b.col + b.cols) - left
    differing = {(r, c) for r, c, _, _ in dataset.diffs}
    location = sheet_location(dataset.sheet, b.address)

    ws.cell(row, 1, f"{result.relative} | {dataset.sheet}!{b.address}").font = TITLE_FONT
    link(ws.cell(row, 2 + cols + 1), "Open in Arco copy", str(result.target), location)
    link(ws.cell(row, 2 + 2 * (cols + 1)), "Open in ResQ original", str(result.source), location)
    ws.cell(row + 1, 1, dataset.arco_formula)
    grids = (("ResQ (saved values)", a), ("Arco (refreshed)", b), ("Difference (Arco - ResQ)", None))
    for grid_index, (title, _) in enumerate(grids):
        first = 2 + grid_index * (cols + 1)
        ws.cell(row + 2, first, title).font = Font(bold=True)
        for c in range(cols):
            ws.cell(row + 3, first + c, get_column_letter(left + c)).font = Font(bold=True, color="7F7F7F")
    for r in range(rows):
        ws.cell(row + 4 + r, 1, top + r).font = Font(bold=True, color="7F7F7F")
        for c in range(cols):
            sheet_row, sheet_col = top + r, left + c
            resq, arco = a.value(sheet_row, sheet_col), b.value(sheet_row, sheet_col)
            diff = arco - resq if is_number(resq) and is_number(arco) else (None if same_value(resq, arco) else "≠")
            for grid_index, value in enumerate((resq, arco, diff)):
                cell = ws.cell(row + 4 + r, 2 + grid_index * (cols + 1) + c, value)
                if is_number(value):
                    cell.number_format = "#,##0.00"
                if (sheet_row, sheet_col) in differing:
                    cell.fill = DIFF_FILL
    return row + 4 + rows + 2


def write_review(books: list[BookResult], path: Path) -> None:
    review = Workbook()
    summary = review.active
    summary.title = "Workbooks"
    datasets_ws = review.create_sheet("Datasets")
    diffs_ws = review.create_sheet("Differences")
    side_ws = review.create_sheet("Side by Side")

    header(summary, ["Workbook", "Datasets", "Different", "Conversion", "Refresh", "Error",
                     "Arco copy", "ResQ original"],
           [50, 10, 10, 60, 60, 50, 14, 14])
    header(datasets_ws, ["Workbook", "Sheet", "Range", "Status", "ResQ size", "Arco size", "Cells",
                         "Different cells", "Max abs diff", "First difference", "Arco formula",
                         "ResQ formula", "Arco copy", "ResQ original", "Side by side"],
           [40, 20, 14, 18, 10, 10, 8, 14, 14, 16, 60, 60, 12, 14, 12])
    header(diffs_ws, ["Workbook", "Sheet", "Cell", "ResQ", "Arco", "Arco - ResQ", "Dataset range",
                      "Arco copy", "ResQ original"],
           [40, 20, 10, 18, 18, 14, 16, 12, 14])
    side_ws.cell(1, 1, f"Cells differing beyond {IMPORT_DECIMAL_PLACES} decimal places are shaded. "
                       "Column A and the grey letters are the cell's row and column in the workbook.")
    side_row = 3

    for result in books:
        differing = [d for d in result.datasets if d.diffs]
        for dataset in differing:
            dataset.side_by_side_row = side_row
            side_row = write_side_by_side(side_ws, result, dataset, side_row)

        row = summary.max_row + 1
        summary.append([str(result.relative), len(result.datasets), len(differing),
                        result.conversion.replace(" Save the workbook to keep the changes.", ""),
                        result.refresh, result.error])
        link(summary.cell(row, 7), "Open", str(result.target), "")
        link(summary.cell(row, 8), "Open", str(result.source), "")
        converted_all = " 0 not updated." in result.conversion
        if result.error or differing or not converted_all or not result.refresh.startswith("Refreshed"):
            summary.cell(row, 1).fill = DIFF_FILL

        for dataset in result.datasets:
            a, b = dataset.resq, dataset.arco
            first = cell_range(dataset.diffs[0][0], dataset.diffs[0][1]) if dataset.diffs else ""
            row = datasets_ws.max_row + 1
            datasets_ws.append([str(result.relative), dataset.sheet, b.address, dataset.status,
                                f"{a.rows} x {a.cols}", f"{b.rows} x {b.cols}",
                                dataset.cells_compared, len(dataset.diffs), dataset.max_abs_diff, first,
                                dataset.arco_formula, dataset.resq_formula])
            datasets_ws.cell(row, 9).number_format = "#,##0.00"
            datasets_ws.cell(row, 4).fill = MATCH_FILL if dataset.status == "Match" else DIFF_FILL
            spot = first or b.address
            link(datasets_ws.cell(row, 13), "Go", str(result.target), sheet_location(dataset.sheet, spot))
            link(datasets_ws.cell(row, 14), "Go", str(result.source), sheet_location(dataset.sheet, spot))
            if dataset.side_by_side_row:
                link(datasets_ws.cell(row, 15), "View", None, f"'Side by Side'!A{dataset.side_by_side_row}")

            for sheet_row, sheet_col, resq, arco in dataset.diffs:
                address = cell_range(sheet_row, sheet_col)
                row = diffs_ws.max_row + 1
                diffs_ws.append([str(result.relative), dataset.sheet, address, resq, arco,
                                 arco - resq if is_number(resq) and is_number(arco) else None, b.address])
                for column in (4, 5, 6):
                    diffs_ws.cell(row, column).number_format = "#,##0.00"
                location = sheet_location(dataset.sheet, address)
                link(diffs_ws.cell(row, 8), "Go", str(result.target), location)
                link(diffs_ws.cell(row, 9), "Go", str(result.source), location)

    for ws in (summary, datasets_ws, diffs_ws):
        ws.auto_filter.ref = ws.dimensions
        for cells in ws.iter_rows(min_row=2):
            for cell in cells:
                cell.alignment = Alignment(vertical="top")
    path.parent.mkdir(parents=True, exist_ok=True)
    review.save(path)


def main() -> int:
    if not INPUT_FOLDER or not OUTPUT_FOLDER:
        print("Set INPUT_FOLDER and OUTPUT_FOLDER at the top of this script first.")
        return 2
    source_root, target_root = Path(INPUT_FOLDER).resolve(), Path(OUTPUT_FOLDER).resolve()
    if source_root == target_root or source_root in target_root.parents:
        print("OUTPUT_FOLDER must be outside INPUT_FOLDER.")
        return 2

    books: list[BookResult] = []
    for source in sorted(source_root.rglob("*")):
        if source.suffix.lower() not in WORKBOOK_SUFFIXES or source.name.startswith("~$"):
            continue
        relative = source.relative_to(source_root)
        target = target_root / relative
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, target)
        os.chmod(target, stat.S_IWRITE | stat.S_IREAD)  # a read-only original stays editable here
        books.append(BookResult(relative, source, target))
    if not books:
        print(f"No workbooks found under {source_root}.")
        return 2
    print(f"Copied {len(books)} workbook(s) to {target_root}.")

    convert_all(books)
    write_review(books, REVIEW_PATH)
    print(f"Review workbook: {REVIEW_PATH}")
    os.startfile(REVIEW_PATH)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
