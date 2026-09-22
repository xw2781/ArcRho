"""Combined ArcRho vs ResQ review for one reserving class: datasets, Result Selections, DFM notes.

Wraps ``dataset_side_by_side_review.py`` (plain triangles/vectors),
``rs_dataset_side_by_side_review.py`` (Result Selection loaded datasets and
Selected Ultimate output) and ``dfm_notes_side_by_side_review.py`` (the
Method Notes on every DFM) so a single target reserving class can be checked
in one run, with one combined workbook, instead of running each script and
opening three separate reports.

Nothing is written back to ArcRho or ResQ.

Run with Python 3.10 from the repository root, on a machine that can reach ResQ:

    py -3.10 python-api/migration/validation/combined_side_by_side_review.py --rc "Legacy\\HOL"

``run_review`` is the same run as one call: it compares, writes the workbook,
and reports what it found. The command line above and the Arco Bridge, which
runs this review for the "Review Reserving Class against ResQ" macro on a
Client PC, both go through it, so neither owns a second version of the run.
"""

from __future__ import annotations

import argparse
import os
import sys
import tempfile
from pathlib import Path

_VALIDATION_DIR = Path(__file__).resolve().parent
if str(_VALIDATION_DIR) not in sys.path:
    sys.path.insert(0, str(_VALIDATION_DIR))

import dataset_side_by_side_review as dsbs  # noqa: E402
import dfm_notes_side_by_side_review as dfmnotes  # noqa: E402
import rs_dataset_side_by_side_review as rssbs  # noqa: E402

# The shape ``run_review`` returns. A Bridge built against one version must not
# be handed a bundle that answers with another.
#
# 2: the DFM Method Notes comparison, which adds its own counts, its own
# reserving-class errors and its own flagged rows to the result.
REVIEW_API_VERSION = 2

# Dataset/Result Selection names to leave out of the review entirely. A name is skipped
# if it CONTAINS any of these substrings (case-insensitive). Add more here as needed.
SKIP_DATASET_NAME_SUBSTRINGS: set[str] = {
    " - May 2026",
    "Growth Adjustment",
    "Accounting Cutoff",
}


def _is_skipped_name(name: str, skip_substrings: set[str]) -> bool:
    folded = name.casefold()
    return any(substring.casefold() in folded for substring in skip_substrings)


def _output_path(project_name: str, rc_paths: list[str]) -> Path:
    label = rc_paths[0].split("\\")[-1] if len(rc_paths) == 1 else f"{len(rc_paths)}rcs"
    safe_label = dsbs._INVALID_SHEET_CHARS.sub("-", label)
    return _VALIDATION_DIR / "results" / f"combined_side_by_side_{project_name}_{safe_label}.xlsx"


def _sheet_title(rc_path: str, suffix: str, used: set[str]) -> str:
    segments = [segment.strip() for segment in rc_path.split("\\") if segment.strip()]
    label = f"{segments[-3]} {segments[-1]}" if len(segments) >= 3 else rc_path
    label = dsbs._INVALID_SHEET_CHARS.sub("-", label)
    base = f"{label} {suffix}"[:31]
    title = base
    suffix_index = 2
    while title.casefold() in used:
        tail = f" ({suffix_index})"
        title = base[: 31 - len(tail)] + tail
        suffix_index += 1
    used.add(title.casefold())
    return title


def write_combined_workbook(
    path: Path,
    *,
    project_name: str,
    rc_paths: list[str],
    source_kinds: tuple[str, ...],
    dataset_records: list[dict],
    dataset_rc_errors: list[tuple[str, str]],
    rs_records: list[dict],
    rs_rc_errors: list[tuple[str, str]],
    notes_records: list[dict],
    notes_rc_errors: list[tuple[str, str]],
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    styles = {
        "bold": Font(bold=True),
        "note": Font(italic=True, color="9C0006"),
        "header_fill": PatternFill("solid", fgColor="DDEBF7"),
        "ultimate_fill": PatternFill("solid", fgColor="FFF2CC"),
        "flag_fill": PatternFill("solid", fgColor="FFC7CE"),
        "center": Alignment(horizontal="center"),
        "link": Font(color="0563C1", underline="single"),
    }

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    used_titles = {"summary"}

    dataset_records_by_rc: dict[str, list[dict]] = {}
    for record in dataset_records:
        dataset_records_by_rc.setdefault(record["rc_path"], []).append(record)
    rs_records_by_rc: dict[str, list[dict]] = {}
    for record in rs_records:
        rs_records_by_rc.setdefault(record["rc_path"], []).append(record)
    notes_records_by_rc: dict[str, list[dict]] = {}
    for record in notes_records:
        notes_records_by_rc.setdefault(record["rc_path"], []).append(record)

    dataset_anchors: dict[tuple[str, str, str], tuple[str, int]] = {}
    rs_anchors: dict[tuple[str, str], tuple[str, int]] = {}
    notes_anchors: dict[tuple[str, str], tuple[str, int]] = {}

    for rc_path in rc_paths:
        ds_records = sorted(
            dataset_records_by_rc.get(rc_path, []), key=lambda r: (r["kind"], r["name"].casefold())
        )
        if ds_records:
            sheet = workbook.create_sheet(_sheet_title(rc_path, "DS", used_titles))
            sheet.cell(row=1, column=1, value=rc_path).font = styles["bold"]
            row = 3
            for record in ds_records:
                header_row, next_row = dsbs._write_dataset_block(sheet, row, record, styles)
                dataset_anchors[(rc_path, record["kind"], record["name"])] = (sheet.title, header_row)
                row = next_row + 2
            dsbs._autosize(sheet, min_width=9, max_width=22)
            sheet.column_dimensions["A"].width = 30
            sheet.freeze_panes = "A3"

        rs_recs = sorted(rs_records_by_rc.get(rc_path, []), key=lambda r: r["name"].casefold())
        if rs_recs:
            sheet = workbook.create_sheet(_sheet_title(rc_path, "RS", used_titles))
            sheet.cell(row=1, column=1, value=rc_path).font = styles["bold"]
            row = 3
            for record in rs_recs:
                header_row, next_row = rssbs._write_rs_block(sheet, row, record, styles)
                rs_anchors[(rc_path, record["name"])] = (sheet.title, header_row)
                row = next_row + 2
            rssbs._autosize(sheet, min_width=9, max_width=22)
            sheet.column_dimensions["A"].width = 26
            sheet.freeze_panes = "A3"

        note_recs = sorted(notes_records_by_rc.get(rc_path, []), key=lambda r: r["name"].casefold())
        if note_recs:
            sheet = workbook.create_sheet(_sheet_title(rc_path, "Notes", used_titles))
            sheet.cell(row=1, column=1, value=rc_path).font = styles["bold"]
            row = 3
            for record in note_recs:
                header_row, next_row = dfmnotes._write_notes_block(sheet, row, record, styles)
                notes_anchors[(rc_path, record["name"])] = (sheet.title, header_row)
                row = next_row + 2
            dfmnotes.size_notes_sheet(sheet)
            sheet.freeze_panes = "A3"

    summary_sheet.cell(
        row=1,
        column=1,
        value=(
            f"Project: {project_name}    Reserving class(es): {', '.join(rc_paths)}    "
            f"Plain-dataset scope: {', '.join(source_kinds)}    "
            "DFM Method Notes are compared line by line, ignoring line endings and trailing blank lines"
        ),
    ).font = styles["bold"]
    headers = ["Type", "RC Path", "Kind", "Name", "Max Abs Diff", "Flagged Cells", "Note"]
    for col, text in enumerate(headers, start=1):
        summary_sheet.cell(row=3, column=col, value=text).font = styles["bold"]
    summary_sheet.freeze_panes = "A4"

    row = 4
    for record in sorted(
        (r for r in dataset_records if r["needs_review"]),
        key=lambda r: (r["rc_path"], r["kind"], r["name"].casefold()),
    ):
        summary_sheet.cell(row=row, column=1, value="Dataset")
        summary_sheet.cell(row=row, column=2, value=record["rc_path"])
        summary_sheet.cell(row=row, column=3, value=record["kind"])
        name_cell = summary_sheet.cell(row=row, column=4, value=record["name"])
        anchor = dataset_anchors.get((record["rc_path"], record["kind"], record["name"]))
        if anchor:
            anchor_sheet_title, anchor_row = anchor
            name_cell.hyperlink = f"#'{anchor_sheet_title}'!A{anchor_row}"
            name_cell.font = styles["link"]
        if record["max_abs_diff"] is not None:
            summary_sheet.cell(row=row, column=5, value=record["max_abs_diff"]).number_format = dsbs.NUMBER_FORMAT
        summary_sheet.cell(row=row, column=6, value=record["flagged_cells"] or None)
        summary_sheet.cell(row=row, column=7, value=record["note"])
        row += 1

    for record in sorted(
        (r for r in rs_records if r["needs_review"]), key=lambda r: (r["rc_path"], r["name"].casefold())
    ):
        summary_sheet.cell(row=row, column=1, value="Result Selection")
        summary_sheet.cell(row=row, column=2, value=record["rc_path"])
        name_cell = summary_sheet.cell(row=row, column=4, value=record["name"])
        anchor = rs_anchors.get((record["rc_path"], record["name"]))
        if anchor:
            anchor_sheet_title, anchor_row = anchor
            name_cell.hyperlink = f"#'{anchor_sheet_title}'!A{anchor_row}"
            name_cell.font = styles["link"]
        if record["max_abs_diff"] is not None:
            summary_sheet.cell(row=row, column=5, value=record["max_abs_diff"]).number_format = rssbs.NUMBER_FORMAT
        summary_sheet.cell(row=row, column=6, value=record["flagged_cells"] or None)
        summary_sheet.cell(row=row, column=7, value=record["note"])
        row += 1

    for record in sorted(
        (r for r in notes_records if r["needs_review"]), key=lambda r: (r["rc_path"], r["name"].casefold())
    ):
        summary_sheet.cell(row=row, column=1, value=dfmnotes.NOTES_TYPE_LABEL)
        summary_sheet.cell(row=row, column=2, value=record["rc_path"])
        summary_sheet.cell(row=row, column=3, value=record["kind"])
        name_cell = summary_sheet.cell(row=row, column=4, value=record["name"])
        anchor = notes_anchors.get((record["rc_path"], record["name"]))
        if anchor:
            anchor_sheet_title, anchor_row = anchor
            name_cell.hyperlink = f"#'{anchor_sheet_title}'!A{anchor_row}"
            name_cell.font = styles["link"]
        # Notes are prose, so there is no difference to measure -- only the
        # count of lines that disagree.
        summary_sheet.cell(row=row, column=6, value=record["differing_lines"] or None)
        summary_sheet.cell(row=row, column=7, value=record["note"])
        row += 1

    for rc_path, note in dataset_rc_errors:
        summary_sheet.cell(row=row, column=1, value="Dataset")
        summary_sheet.cell(row=row, column=2, value=rc_path)
        summary_sheet.cell(row=row, column=3, value="(reserving class)")
        summary_sheet.cell(row=row, column=7, value=note)
        row += 1
    for rc_path, note in rs_rc_errors:
        summary_sheet.cell(row=row, column=1, value="Result Selection")
        summary_sheet.cell(row=row, column=2, value=rc_path)
        summary_sheet.cell(row=row, column=3, value="(reserving class)")
        summary_sheet.cell(row=row, column=7, value=note)
        row += 1
    for rc_path, note in notes_rc_errors:
        summary_sheet.cell(row=row, column=1, value=dfmnotes.NOTES_TYPE_LABEL)
        summary_sheet.cell(row=row, column=2, value=rc_path)
        summary_sheet.cell(row=row, column=3, value="(reserving class)")
        summary_sheet.cell(row=row, column=7, value=note)
        row += 1

    if row == 4:
        summary_sheet.cell(row=row, column=1, value="Nothing needs review.")

    dsbs._autosize(summary_sheet, min_width=12, max_width=80)

    descriptor, temporary_name = tempfile.mkstemp(prefix=f".{path.stem}-", suffix=".xlsx", dir=path.parent)
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, path)
    finally:
        workbook.close()
        if temporary_path.exists():
            temporary_path.unlink()


def _flagged_rows(
    dataset_records: list[dict], rs_records: list[dict], notes_records: list[dict]
) -> list[dict]:
    """Every block the workbook's Summary links to, as plain JSON values.

    The caller that cannot open the workbook -- a macro reporting the run in a
    dialog -- shows this instead, so what needs attention is named in both
    places from one list. A notes row carries no difference to measure, so its
    count of disagreeing lines stands in for the flagged-cell count.
    """

    rows: list[dict] = []
    for kind_label, records, sort_key in (
        ("Dataset", dataset_records, lambda r: (r["rc_path"], r["kind"], r["name"].casefold())),
        ("Result Selection", rs_records, lambda r: (r["rc_path"], r["name"].casefold())),
        (dfmnotes.NOTES_TYPE_LABEL, notes_records, lambda r: (r["rc_path"], r["name"].casefold())),
    ):
        for record in sorted((r for r in records if r["needs_review"]), key=sort_key):
            rows.append({
                "type": kind_label,
                "rc_path": record["rc_path"],
                "kind": str(record.get("kind") or kind_label),
                "name": record["name"],
                "max_abs_diff": record.get("max_abs_diff"),
                "flagged_cells": record.get("flagged_cells", record.get("differing_lines", 0)),
                "note": record["note"],
            })
    return rows


def run_review(
    *,
    project_name: str,
    rc_paths: list[str],
    source_kinds: tuple[str, ...] = dsbs.SOURCE_KINDS,
    skip_substrings: set[str] | None = None,
    output_path: Path | None = None,
    credentials: dict[str, str] | None = None,
    progress=print,
) -> dict:
    """Compare one or more reserving classes, write the workbook, and report.

    This is the whole run: the command line below and the Arco Bridge both
    call it, so the scope rules, the skip list, and the workbook are defined
    once. Three comparisons make it up -- plain datasets, Result Selections
    and the Method Notes on every DFM. Nothing is written back to ArcRho or
    ResQ.
    """

    rc_paths = [str(path).strip() for path in rc_paths if str(path).strip()]
    if not rc_paths:
        raise ValueError("At least one reserving-class path is required.")
    skipped = SKIP_DATASET_NAME_SUBSTRINGS if skip_substrings is None else set(skip_substrings)

    progress("Comparing plain datasets...")
    dataset_records, dataset_rc_errors = dsbs.run_comparison(
        project_name=project_name,
        rc_paths=rc_paths,
        source_kinds=source_kinds,
        credentials=credentials,
        progress=progress,
    )

    progress("Comparing Result Selections...")
    rs_records, rs_rc_errors = rssbs.run_comparison(
        project_name=project_name,
        rc_paths=rc_paths,
        credentials=credentials,
        progress=progress,
    )

    progress("Comparing DFM Method Notes...")
    notes_records, notes_rc_errors = dfmnotes.run_comparison(
        project_name=project_name,
        rc_paths=rc_paths,
        credentials=credentials,
        progress=progress,
    )

    skipped_datasets = 0
    skipped_result_selections = 0
    skipped_dfm_notes = 0
    if skipped:
        skipped_datasets = sum(1 for r in dataset_records if _is_skipped_name(r["name"], skipped))
        skipped_result_selections = sum(1 for r in rs_records if _is_skipped_name(r["name"], skipped))
        skipped_dfm_notes = sum(1 for r in notes_records if _is_skipped_name(r["name"], skipped))
        dataset_records = [r for r in dataset_records if not _is_skipped_name(r["name"], skipped)]
        rs_records = [r for r in rs_records if not _is_skipped_name(r["name"], skipped)]
        notes_records = [r for r in notes_records if not _is_skipped_name(r["name"], skipped)]
        progress(
            f"Skipping {skipped_datasets} dataset(s), "
            f"{skipped_result_selections} Result Selection(s) and "
            f"{skipped_dfm_notes} DFM(s) by name."
        )

    workbook_path = Path(output_path) if output_path is not None else _output_path(project_name, rc_paths)
    write_combined_workbook(
        workbook_path,
        project_name=project_name,
        rc_paths=rc_paths,
        source_kinds=source_kinds,
        dataset_records=dataset_records,
        dataset_rc_errors=dataset_rc_errors,
        rs_records=rs_records,
        rs_rc_errors=rs_rc_errors,
        notes_records=notes_records,
        notes_rc_errors=notes_rc_errors,
    )

    dataset_needs_review = [r for r in dataset_records if r["needs_review"]]
    rs_needs_review = [r for r in rs_records if r["needs_review"]]
    notes_needs_review = [r for r in notes_records if r["needs_review"]]
    rc_errors = [
        {"rc_path": rc_path, "type": label, "note": note}
        for label, entries in (
            ("Dataset", dataset_rc_errors),
            ("Result Selection", rs_rc_errors),
            (dfmnotes.NOTES_TYPE_LABEL, notes_rc_errors),
        )
        for rc_path, note in entries
    ]
    return {
        "review_api_version": REVIEW_API_VERSION,
        "project_name": project_name,
        "rc_paths": list(rc_paths),
        "source_kinds": list(source_kinds),
        "workbook_path": str(workbook_path),
        "workbook_name": workbook_path.name,
        "datasets_compared": len(dataset_records),
        "datasets_needing_review": len(dataset_needs_review),
        "result_selections_compared": len(rs_records),
        "result_selections_needing_review": len(rs_needs_review),
        "dfm_notes_compared": len(notes_records),
        "dfm_notes_needing_review": len(notes_needs_review),
        "skipped_datasets": skipped_datasets,
        "skipped_result_selections": skipped_result_selections,
        "skipped_dfm_notes": skipped_dfm_notes,
        "reserving_class_errors": rc_errors,
        "flagged": _flagged_rows(dataset_records, rs_records, notes_records),
        "needs_attention": bool(
            dataset_needs_review or rs_needs_review or notes_needs_review or rc_errors
        ),
    }


def _parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--project", default=dsbs.TARGET_PROJECT_NAME, help="ResQ project name to review.")
    parser.add_argument(
        "--rc",
        action="append",
        required=True,
        help="Only review reserving classes whose path contains this text; repeatable.",
    )
    parser.add_argument(
        "--source-kind",
        action="append",
        choices=[*dsbs.SOURCE_KINDS, "all"],
        help="Plain-dataset source kind to include; repeatable. Defaults to all kinds for a targeted RC check.",
    )
    parser.add_argument("--no-open", action="store_true", help="Do not open the workbook when the run finishes.")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = _parse_args(argv)

    needles = [text.casefold() for text in args.rc]
    rc_paths = [path for path in dsbs.RC_PATHS if any(needle in path.casefold() for needle in needles)]
    if not rc_paths:
        print("No reserving class matched --rc.")
        return 2

    selected = args.source_kind or ["all"]
    source_kinds = dsbs.SOURCE_KINDS if "all" in selected else tuple(dict.fromkeys(selected))

    print(f"Reserving class(es): {', '.join(rc_paths)}")
    result = run_review(
        project_name=args.project,
        rc_paths=rc_paths,
        source_kinds=source_kinds,
    )

    output_path = Path(result["workbook_path"])
    print(
        f"Plain datasets: {result['datasets_compared']} compared, "
        f"{result['datasets_needing_review']} need review."
    )
    print(
        f"Result Selections: {result['result_selections_compared']} compared, "
        f"{result['result_selections_needing_review']} need review."
    )
    print(
        f"DFM Method Notes: {result['dfm_notes_compared']} compared, "
        f"{result['dfm_notes_needing_review']} need review."
    )
    print(f"Excel report: {output_path}")
    if result["needs_attention"] and not args.no_open:
        try:
            os.startfile(output_path)  # noqa: S606 - opening the report just written, for the operator running this script
        except Exception as exc:
            print(f"Could not open the report automatically: {type(exc).__name__}: {exc}")
    return 0 if not result["needs_attention"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
